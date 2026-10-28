"""stdlib unittest suite. Run: python -m unittest discover -s tests -v
CWD must be the project root (discover run from there, as in README).
"""
import math
import os
import sys
import unittest

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from src import flash as flash_mod  # noqa: E402
from src import flux as flux_mod  # noqa: E402
from src import forecast as forecast_mod  # noqa: E402
from src import dashboard as dashboard_mod  # noqa: E402
from pptx import Presentation  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class TestFluxGroundTruth(unittest.TestCase):
    def test_injected_events_recovered(self):
        # Volume spike Jun 2024 -> commissions explained as Volume.
        rep = flux_mod.flux_report("2024-06")
        comm = [r for r in rep["lines"] if r["line"] == "commissions"][0]
        self.assertTrue(comm["material"])
        self.assertIn("Volume", comm["explanation"])
        # Fee-schedule change Mar 2025 -> platform_fees explained as Price.
        rep = flux_mod.flux_report("2025-03")
        plat = [r for r in rep["lines"] if r["line"] == "platform_fees"][0]
        self.assertTrue(plat["material"])
        self.assertIn("Price", plat["explanation"])
        # One-off legal expense Jul 2025 -> one_offs explained as One-off.
        rep = flux_mod.flux_report("2025-07")
        one = [r for r in rep["lines"] if r["line"] == "one_offs"][0]
        self.assertTrue(one["material"])
        self.assertIn("One-off", one["explanation"])
        # Client onboarding Sep 2025 -> client note names C041 onboarding.
        rep = flux_mod.flux_report("2025-09")
        self.assertIn("C041", rep["client_note"])
        self.assertIn("onboarding", rep["client_note"].lower())


class TestFlashReconciles(unittest.TestCase):
    def test_flash_totals_reconcile(self):
        df = pd.read_csv(os.path.join(ROOT, "data/synthetic/daily_flash.csv"))
        rev_lines = ["commissions", "payment_for_order_flow", "margin_interest",
                     "platform_fees", "custody_fees"]
        cost_lines = ["clearing", "compensation", "technology", "occupancy", "one_offs"]
        for _, row in df.iterrows():
            self.assertAlmostEqual(row[rev_lines].sum(), row["total_revenue"], delta=0.02)
            self.assertAlmostEqual(row[cost_lines].sum(), row["total_costs"], delta=0.02)
            self.assertAlmostEqual(row["total_revenue"] - row["total_costs"],
                                   row["net_pnl"], delta=0.02)
        rep = flash_mod.flash_report("2025-12-30")
        self.assertAlmostEqual(rep["row"]["net_pnl"],
                               rep["row"]["total_revenue"] - rep["row"]["total_costs"],
                               delta=0.02)


class TestForecast(unittest.TestCase):
    def test_forecast_returns_finite_metrics(self):
        cmp_df = forecast_mod.compare()
        self.assertEqual(len(cmp_df), 6)  # 5 lines + total
        for col in ["ols_mape", "naive_mape", "ols_rmse", "naive_rmse"]:
            self.assertTrue(cmp_df[col].apply(math.isfinite).all())
            self.assertTrue((cmp_df[col] >= 0).all())
        flags = forecast_mod.flag_anomalies()
        self.assertIsInstance(flags, pd.DataFrame)


class TestDashboard(unittest.TestCase):
    def test_dashboard_build(self):
        path = dashboard_mod.build()
        self.assertTrue(os.path.exists(path))
        html = open(path).read()
        self.assertGreater(len(html), 10_000)
        # KPI values must appear in the file.
        self.assertIn("Latest month (2025-12) revenue:", html)
        self.assertIn("$2,663,624", html)


class TestDeck(unittest.TestCase):
    def test_deck_has_10_slides(self):
        path = os.path.join(ROOT, "reports/DriveWealth_FPA_Consultant_Deck.pptx")
        prs = Presentation(path)
        self.assertEqual(len(prs.slides._sldIdLst), 10)
        titles = [s.shapes.title.text for s in prs.slides]
        self.assertEqual(titles[0][:4], "FP&A")
        self.assertIn("About This Project", titles[-1])


if __name__ == "__main__":
    unittest.main()
