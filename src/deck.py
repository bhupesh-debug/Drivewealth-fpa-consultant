"""Builds the 10-slide board-style deck (python-pptx, real content).

Slide order (jd-to-ml-project structure):
1 Title, 2 Problem Overview, 3 JD Breakdown, 4 Approach, 5 Data Strategy,
6 Model Design, 7 Results & Insights, 8 Business Impact,
9 Future Improvements, 10 About This Project.
CLI: python -m src.deck
"""
import os
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data", "synthetic")
REPORTS = os.path.join(ROOT, "reports")


def add_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(16)
        p.space_after = Pt(6)
    return slide


def build():
    from src.forecast import compare
    monthly = pd.read_csv(os.path.join(DATA, "monthly_close.csv"))
    latest = monthly.month.max()
    rev = monthly[(monthly.month == latest) & (monthly.category == "revenue")]["actual"].sum()
    cost = monthly[(monthly.month == latest) & (monthly.category == "cost")]["actual"].sum()
    cmp_df = compare()
    tot = cmp_df[cmp_df.line == "total_revenue"].iloc[0]
    beats = int(cmp_df["ols_beats_naive_mape"].sum())
    nlines = len(cmp_df)
    comp_lines = []
    for _, r in cmp_df.iterrows():
        comp_lines.append(f"{r['line']}: OLS MAPE {r['ols_mape']:.2f}% vs "
                          f"naive {r['naive_mape']:.2f}%")
    prs = Presentation()
    prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
    add_slide(prs, "FP&A Analytics for a Brokerage: Monthly Close, Daily Flash & Client Revenue",
              ["A rules-based FP&A analytics portfolio project — built for an FP&A Consultant role",
               "Fictional firm: NorthPeak Securities. ALL DATA SYNTHETIC — no real company data.",
               "Author: Bhupesh Chandra Dimri — bhupesh7cd@gmail.com"])
    add_slide(prs, "Problem Overview",
              ["Monthly close flux analysis is slow and narrative-heavy when done by hand.",
               "Daily P&L flash reporting needs same-day movers and flags, not just totals.",
               "Client revenue dashboards must tie segment and client moves to the P&L.",
               "Board and exec reporting needs one consistent, reproducible set of numbers."])
    add_slide(prs, "JD Breakdown — the five focuses",
              ["1. Monthly close flux analysis — variance vs prior month and vs budget.",
               "2. Daily P&L flash reporting — day vs prior day vs month-to-date.",
               "3. Client revenue dashboards — segment and top-client views.",
               "4. Board decks — this 10-slide deck, built programmatically.",
               "5. Exec reporting — one-page exec summary with verified numbers.",
               "Tools in scope: Excel-style close tables, PowerPoint, Databricks-style "
               "pipelines (replicated here in Python), AI-assisted drafting."])
    add_slide(prs, "Approach",
              ["Honest-fit design: this is a PURE FP&A role with no ML requirement.",
               "So the core is rules-based: documented materiality thresholds and driver "
               "attribution (volume / price / mix / one-off / client concentration).",
               "Forecasting uses transparent numpy OLS trend+seasonal, benchmarked "
               "against a naive seasonal baseline — no black-box model, no LLM claims.",
               "A rules-based explainer beating a black box on this data is a feature: "
               "it is auditable, explainable to auditors and executives."])
    add_slide(prs, "Data Strategy",
              ["100% synthetic, seeded and reproducible (python -m src.data_gen --seed 42).",
               "24 months monthly close (Jan 2024–Dec 2025): actuals + budget, 5 revenue "
               "lines, 5 cost lines shaped like a brokerage P&L.",
               "42 fictional clients (retail / institutional / introducing brokers) with "
               "monthly revenue; 90 days of daily P&L flash data.",
               "Four KNOWN injected flux events with recorded ground truth, so explainer "
               "accuracy is measurable — never presented as real events."])
    add_slide(prs, "Model Design",
              ["Flux engine: material if |variance| >= $25,000 AND >= 5% vs prior month; "
               "one-offs also flagged above 3x their 24-month median.",
               "Attribution priority: one-off spike, then volume (commissions), then fee "
               "schedule (platform fees), then client onboarding/offboarding, else mix/other.",
               "Flash: day vs prior day vs MTD, top movers, flags at >=20% line moves and "
               "|z| > 2 on net P&L; anomaly flagger at |z| > 3 vs prior 30 days.",
               "Forecast: OLS [1, trend, sin/cos seasonality] vs naive same-month-last-year."])
    add_slide(prs, "Results & Insights",
              [f"Latest synthetic month ({latest}): revenue ${rev:,.0f}, costs "
               f"${cost:,.0f}, net P&L ${rev - cost:,.0f}.",
               f"Forecast hold-out (Oct–Dec 2025), total revenue: OLS MAPE "
               f"{tot['ols_mape']:.2f}% vs naive seasonal {tot['naive_mape']:.2f}%.",
               f"OLS beat the naive baseline on MAPE for {beats} of {nlines} lines "
               "(see report for the honest per-line table — baselines win some lines)."]
              + comp_lines)
    add_slide(prs, "Business Impact",
              ["Close narrative in minutes: every material line gets a plain-English "
               "driver explanation instead of a blank variance table.",
               "Flash discipline: same-day movers and flags surface what changed before "
               "the close, when it can still be investigated.",
               "Client lens: segment and top-client views tie revenue moves to the "
               "clients behind them, including onboarding/offboarding.",
               "Reproducibility: one seed regenerates every figure in this deck."])
    add_slide(prs, "Future Improvements",
              ["Replace synthetic data with real close data under proper access controls; "
               "re-validate thresholds on real variance distributions.",
               "Add Databricks / SQL pipeline version of the flux and flash engines.",
               "Add Excel workbook export of the flux report for close reviewers.",
               "Extend forecasting to client-level models once real client history exists; "
               "keep the naive baseline as the honesty check."])
    add_slide(prs, "About This Project",
              ["Author: Bhupesh Chandra Dimri",
               "Contact: bhupesh7cd@gmail.com",
               "GitHub: link placeholder — repository to be published by the author.",
               "Built with Python (pandas, numpy, matplotlib, python-pptx, reportlab); "
               "no sklearn, no Streamlit, no Databricks in this build — stated openly.",
               "All data synthetic; fictional firm NorthPeak Securities."])
    assert len(prs.slides.__iter__.__self__._sldIdLst) == 10, "deck must be 10 slides"
    os.makedirs(REPORTS, exist_ok=True)
    path = os.path.join(REPORTS, "DriveWealth_FPA_Consultant_Deck.pptx")
    prs.save(path)
    print(f"Written: {path} ({len(prs.slides._sldIdLst)} slides)")
    return path


if __name__ == "__main__":
    build()
